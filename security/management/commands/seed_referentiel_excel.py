import re
import openpyxl
from django.core.management.base import BaseCommand
from referentiel.models import TypeReferentiel, Reference, ReferentielItem, Region
from user.models import EntiteMetier

REGIONS_VALIDES = ['DRY', 'DRC', 'DRSANO', 'DRSOM', 'DRE', 'DRONO', 'DRSM', 'DRNEA', 'DRD']
ALIAS_POSTE = {'NYOM II': 'NYOM', 'MILE 2 LIMBE': 'MILE2 LIMBE'}


def normaliser_poste(poste):
    return ALIAS_POSTE.get(poste, poste)


def extraire_poste(segment, ouvrage):
    if segment.startswith('PRODUCTION'):
        parts = segment.split(' - ')
        return normaliser_poste(parts[-1].strip() if len(parts) > 1 else segment.strip())
    match = re.match(r'^([^_]+(?:\s[^_]+)*?)(?:_POSTE SOURCE|_LIGNE|$)', ouvrage.strip())
    poste = match.group(1).strip() if match else ouvrage.split('_')[0].strip()
    return normaliser_poste(poste)


def extraire_type_poste(segment, ouvrage):
    u = ouvrage.upper()
    if 'POSTE SOURCE' in u or re.search(r'POSTE DE \w+', u):
        return 'POSTE_SOURCE'
    if 'LIGNE' in u:
        return 'LIGNE'
    if segment.startswith('PRODUCTION'):
        return 'PRODUCTION'
    return 'DEPART'


def extraire_region_code(segment):
    if not segment.startswith('DISTRIBUTION-'):
        return None
    code = segment.replace('DISTRIBUTION-', '').strip()
    return code if code in REGIONS_VALIDES else None


def extraire_numero_rame(texte):
    """'TRANSFO N°1 90/15kV' ou 'RAME 15kV N°1' -> '1'"""
    match = re.search(r'N\s*[°ºo]\s*(\d+)', texte, re.IGNORECASE)
    return match.group(1) if match else None


class Command(BaseCommand):
    help = 'Génère le référentiel depuis le fichier Excel BD_asset'

    def add_arguments(self, parser):
        parser.add_argument('excel_path', type=str)

    def handle(self, *args, **options):
        wb = openpyxl.load_workbook(options['excel_path'], data_only=True)

        types = {}
        for nom in ['SEGMENT', 'POSTE', 'TYPE_POSTE', 'RAME', 'OUVRAGE', 'DEPART']:
            types[nom], _ = TypeReferentiel.objects.get_or_create(nom=nom)

        regions = {}
        for code in REGIONS_VALIDES:
            regions[code], _ = Region.objects.get_or_create(code=code)

        entites = {
            'PRODUCTION': EntiteMetier.objects.get(type='PROD'),
            'TRANSPORT': EntiteMetier.objects.get(type='TRANS'),
            'DISTRIBUTION': EntiteMetier.objects.get(type='DIST'),
        }

        sheets_config = [
            ('Type réseau et Réf production', 'PRODUCTION', 3),
            ('Type de réseau et réf transport', 'TRANSPORT', 4),
            ('Type de réseau et réf distribut', 'DISTRIBUTION', 5),
        ]

        # PASSE 1 : poste -> région
        poste_region = {}
        ws_dist = wb['Type de réseau et réf distribut']
        for row in ws_dist.iter_rows(min_row=3, values_only=True):
            if not row[0]:
                continue
            segment = str(row[0]).strip()
            ouvrage = str(row[1]).strip() if row[1] else ''
            region_code = extraire_region_code(segment)
            if region_code:
                poste_region[extraire_poste(segment, ouvrage)] = region_code

        # PASSE 2 : créer Reference + items
        total = 0
        sans_region = set()

        for sheet_name, entite_key, ref_col_index in sheets_config:
            ws = wb[sheet_name]
            for row in ws.iter_rows(min_row=3, values_only=True):
                if not row[0]:
                    continue

                segment = str(row[0]).strip()
                ouvrage = str(row[1]).strip() if row[1] else ''
                gr_tfo = str(row[2]).strip() if row[2] else ''
                depart = str(row[3]).strip() if entite_key == 'DISTRIBUTION' and row[3] else None
                if entite_key == 'PRODUCTION':
                    reference_val = '_'.join(p for p in [segment, ouvrage, gr_tfo] if p)
                else:
                    reference_val = str(row[ref_col_index - 1]).strip() if row[ref_col_index - 1] else ''

                if not reference_val or reference_val == 'None':
                    continue

                poste = extraire_poste(segment, ouvrage)
                type_poste = extraire_type_poste(segment, ouvrage)
                region_code = poste_region.get(poste)
                if not region_code:
                    sans_region.add((poste, entite_key))

                ref, _ = Reference.objects.get_or_create(
                    valeur=reference_val,
                    defaults={
                        "entite_metier": entites[entite_key],
                        "region": regions.get(region_code) if region_code else None
                    }
                )
                if region_code and not ref.region_id: # type: ignore
                    ref.region = regions[region_code] 
                    ref.save(update_fields=['region'])

                items_to_create = [('SEGMENT', segment), ('POSTE', poste), ('TYPE_POSTE', type_poste)]
                if gr_tfo:
                    items_to_create.append(('OUVRAGE', gr_tfo))
                if depart:
                    items_to_create.append(('DEPART', depart))

                # Extraction RAME — uniquement côté DISTRIBUTION
                if entite_key == 'DISTRIBUTION':
                    if segment == 'DISTRIBUTION-MAINTENANCE POSTES' and depart:
                        # ici "depart" (row[3]) contient en réalité "RAME 15kV N°1"
                        numero = extraire_numero_rame(depart)
                        if numero:
                            items_to_create = [it for it in items_to_create if it[0] != 'DEPART']
                            items_to_create.append(('RAME', f"{poste}_RAME{numero}"))
                    else:
                        numero = extraire_numero_rame(gr_tfo)
                        if numero:
                            items_to_create.append(('RAME', f"{poste}_RAME{numero}"))

                for type_nom, valeur in items_to_create:
                    ReferentielItem.objects.get_or_create(
                        reference=ref, type=types[type_nom], defaults={"valeur": valeur}
                    )
                total += 1

        self.stdout.write(self.style.SUCCESS(f"✅ {total} références traitées"))
        if sans_region:
            self.stdout.write(self.style.WARNING(f"⚠️  {len(sans_region)} poste(s) sans région :"))
            for poste, entite in sorted(sans_region):
                self.stdout.write(f"     - {poste} ({entite})")